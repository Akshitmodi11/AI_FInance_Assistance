"""Word (.docx) loader.

Concepts this teaches:
* A .docx is a zip of XML. ``python-docx`` exposes it as a list of paragraphs, each with a
  *style* ("Normal", "Heading 1", "Heading 2", "List Bullet", ...).
* Headings are structure you get for free — split on them, exactly like the Markdown loader
  splits on ``##``. (Text that is merely bold is NOT a heading; that's why the authoring
  guide says to use real heading styles.)

Library: ``import docx`` → ``docx.Document(path).paragraphs`` → ``p.text``, ``p.style.name``.
Note the name clash: our model is also called ``Document`` — import the library as a module.
"""

from __future__ import annotations

from pathlib import Path

import docx

from src.rag.models import Document, SourceMeta


def load_docx(path: Path, meta: dict | None = None) -> list[Document]:
    """Load a Word document into one Document per heading section.

    Requirements (see tests/unit/test_loaders.py):
    * ``meta`` from manifest.yaml is required (``ValueError`` if None).
    * Any paragraph whose style name starts with ``"Heading"`` starts a new section;
      its text becomes ``SourceMeta.section``. Text before the first heading goes in a
      section named after the title.
    * ``file_type="docx"``, ``doc_id = f"{path.stem}:s{i}"``.
    * Join paragraph texts with ``"\\n"``; skip empty paragraphs and empty sections.
    * Bonus: include table cell text (``document.tables``) — finance guides love tables.
    """
    if meta is None:
        raise ValueError(f"{path.name}: a Word file needs a manifest.yaml entry to be citable")

    base = SourceMeta(**meta, file_type="docx")
    sections: list[tuple[str, list[str]]] = [(base.title, [])]

    for para in docx.Document(path).paragraphs:
        text = para.text.strip()
        if para.style.name.startswith("Heading"):
            sections.append((text, []))
        elif text:
            sections[-1][1].append(text)

    docs: list[Document] = []
    for name, lines in sections:
        if not lines:
            continue  # e.g. no text before the first heading
        docs.append(
            Document(
                doc_id=f"{path.stem}:s{len(docs)}",
                text="\n".join(lines),
                meta=base.model_copy(update={"section": name}),
            )
        )
    return docs
