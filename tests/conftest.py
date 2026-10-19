"""Shared test fixtures.

``FakeEmbedder`` lets RAG tests run in milliseconds without downloading the real model.
It hashes words into buckets, so texts sharing words get similar vectors — enough to test
plumbing (ordering, filtering, fusion), not semantic quality. Semantic quality is measured
by the retrieval eval against the real model.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

import numpy as np
import pytest

from src.rag.models import Chunk, Document, SourceMeta


class FakeEmbedder:
    dim = 64

    def embed(self, texts: list[str]) -> np.ndarray:
        out = np.zeros((len(texts), self.dim), dtype=np.float32)
        for row, text in enumerate(texts):
            for word in re.findall(r"[a-z0-9]+", text.lower()):
                bucket = int(hashlib.md5(word.encode()).hexdigest(), 16) % self.dim
                out[row, bucket] += 1.0
            norm = np.linalg.norm(out[row])
            if norm:
                out[row] /= norm
        return out


@pytest.fixture
def fake_embedder() -> FakeEmbedder:
    return FakeEmbedder()


def make_meta(**overrides) -> SourceMeta:
    base = dict(
        title="Test Article",
        url="https://example.com/a",
        source="Example",
        category="basics",
        file_type="md",
    )
    return SourceMeta(**{**base, **overrides})


def make_doc(text: str, doc_id: str = "doc:s0", **meta) -> Document:
    return Document(doc_id=doc_id, text=text, meta=make_meta(**meta))


def make_chunk(text: str, chunk_id: str, **meta) -> Chunk:
    doc_id = chunk_id.rsplit(":", 1)[0]
    return Chunk(chunk_id=chunk_id, doc_id=doc_id, text=text, meta=make_meta(**meta))


def make_pdf(pages: list[list[str]]) -> bytes:
    """Build a minimal valid PDF with one text line per entry. No PDF library needed."""
    n = len(pages)
    page_ids = [4 + 2 * i for i in range(n)]
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {n} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, lines in enumerate(pages):
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_ids[i] + 1} 0 R >>".encode()
        )
        ops, y = [], 740
        for line in lines:
            esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"BT /F1 11 Tf 72 {y} Td ({esc}) Tj ET")
            y -= 16
        stream = "\n".join(ops).encode("latin-1")
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(path: Path, blocks: list[tuple[str, str]]) -> Path:
    """Write a .docx. ``blocks`` = [(style, text)], style like "Heading 1" or "Normal"."""
    import docx

    document = docx.Document()
    for style, text in blocks:
        document.add_paragraph(text, style=style)
    document.save(path)
    return path


MANIFEST_META = {
    "title": "IRS Publication 590-A",
    "url": "https://www.irs.gov/pub/irs-pdf/p590a.pdf",
    "source": "IRS",
    "category": "retirement_tax",
    "tax_year": 2025,
}


@pytest.fixture
def manifest_meta() -> dict:
    return dict(MANIFEST_META)
